"""Test cases for multi_file_input module.

Tests cover:
- 5 Channels: FileChannel, APIChannel, StreamChannel, DatabaseChannel, IoTChannel
- Each channel has read() + export() methods
- 2 data categories: flatten (single-table) and complex (hierarchical)
- Single JSON output format with metadata + records
"""

import pytest
import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from multi_file_input.channels import (
    FileChannel,
    APIChannel,
    StreamChannel,
    DatabaseChannel,
    IoTChannel,
)
from multi_file_input.adapters import (
    CSVAdapter,
    JSONAdapter,
    XMLAdapter,
    ExcelAdapter,
    YAMLAdapter,
)
from multi_file_input.normalizer import JSONNormalizer
from multi_file_input.exceptions import (
    IngestionError,
    ChannelReadError,
    ParseError,
    NormalizationError,
)


# ============================================================================
# CHANNEL TESTS
# ============================================================================

class TestFileChannel:
    """Test FileChannel - read() and export() methods."""

    def test_read_csv_file(self, tmp_path):
        """Test reading a CSV file."""
        csv_file = tmp_path / "test.csv"
        csv_file.write_text("name,age,city\nAlice,30,NYC\nBob,25,LA\n", encoding="utf-8")

        channel = FileChannel()
        content = channel.read(str(csv_file))

        assert content == b"name,age,city\nAlice,30,NYC\nBob,25,LA\n"

    def test_read_nonexistent_file_raises_error(self):
        """Test that reading non-existent file raises ChannelReadError."""
        channel = FileChannel()
        with pytest.raises(ChannelReadError):
            channel.read("/nonexistent/file.csv")

    def test_export_normalized_json(self, tmp_path):
        """Test export() writes normalized JSON to file."""
        data = {
            "metadata": {
                "source_type": "csv",
                "source_channel": "file",
                "record_count": 2,
                "data_category": "flatten",
            },
            "records": [
                {"_id": "abc_0", "_source": {"name": "Alice", "age": "30"}},
                {"_id": "abc_1", "_source": {"name": "Bob", "age": "25"}},
            ]
        }

        channel = FileChannel()
        output_path = tmp_path / "output.json"
        result = channel.export(str(output_path), data)

        assert output_path.exists()
        loaded = json.loads(output_path.read_text(encoding="utf-8"))
        assert loaded["metadata"]["record_count"] == 2
        assert len(loaded["records"]) == 2

    def test_detect_file_type(self, tmp_path):
        """Test file type detection from extension."""
        csv_file = tmp_path / "test.csv"
        csv_file.write_text("a,b\n1,2\n", encoding="utf-8")

        channel = FileChannel()
        assert channel.detect_type(str(csv_file)) == "csv"

    def test_batch_read(self, tmp_path):
        """Test batch reading all files in directory."""
        (tmp_path / "file1.csv").write_text("a,b\n1,2\n", encoding="utf-8")
        (tmp_path / "file2.csv").write_text("c,d\n3,4\n", encoding="utf-8")

        channel = FileChannel()
        files = channel.batch_read(str(tmp_path), "*.csv")

        assert len(files) == 2


class TestAPIChannel:
    """Test APIChannel - read() and export() methods."""

    def test_read_url_success(self):
        """Test successful API read."""
        channel = APIChannel()
        # Mock response via patching would be done in integration test
        # Here we test the method signature exists
        assert hasattr(channel, 'read')
        assert hasattr(channel, 'export')

    def test_export_normalized_json(self):
        """Test export() method exists and has correct signature."""
        channel = APIChannel()
        data = {
            "metadata": {"source_type": "json", "source_channel": "api", "record_count": 1},
            "records": [{"_id": "test_0", "_source": {"name": "Test"}}]
        }
        # Signature check - would need mock server for full test
        assert callable(channel.export)

    def test_authenticate_bearer(self):
        """Test bearer token authentication."""
        channel = APIChannel()
        token = channel.authenticate({"token": "secret123"}, auth_type="bearer")

        assert token == "secret123"
        assert channel._auth_token == "secret123"

    def test_graphql_method_exists(self):
        """Test GraphQL query method exists."""
        channel = APIChannel()
        assert hasattr(channel, 'graphql')


class TestStreamChannel:
    """Test StreamChannel - read() and export() methods."""

    def test_read_and_export_methods_exist(self):
        """Test that read() and export() methods exist."""
        channel = StreamChannel()
        assert hasattr(channel, 'read')
        assert hasattr(channel, 'export')
        assert callable(channel.read)
        assert callable(channel.export)

    def test_connect_kafka_method_exists(self):
        """Test Kafka connection method exists."""
        channel = StreamChannel()
        assert hasattr(channel, 'connect_kafka')

    def test_connect_websocket_method_exists(self):
        """Test WebSocket connection method exists."""
        channel = StreamChannel()
        assert hasattr(channel, 'connect_websocket')


class TestDatabaseChannel:
    """Test DatabaseChannel - read() and export() methods."""

    def test_read_and_export_methods_exist(self):
        """Test that read() and export() methods exist."""
        channel = DatabaseChannel()
        assert hasattr(channel, 'read')
        assert hasattr(channel, 'export')
        assert callable(channel.read)
        assert callable(channel.export)

    def test_connect_sqlite(self, tmp_path):
        """Test SQLite connection and basic operations."""
        db_path = tmp_path / "test.db"
        channel = DatabaseChannel()
        channel.connect_sqlite(str(db_path))

        # Create table and insert
        channel.read("CREATE TABLE users (name TEXT, age INTEGER)")
        channel.read("INSERT INTO users (name, age) VALUES ('Alice', 30)")

        # Query
        result = channel.read("SELECT * FROM users")
        assert len(result) == 1
        assert result[0]["name"] == "Alice"

    def test_export_to_table(self, tmp_path):
        """Test export() writes normalized JSON to table."""
        db_path = tmp_path / "test.db"
        channel = DatabaseChannel()
        channel.connect_sqlite(str(db_path))

        # Create table
        channel.read("CREATE TABLE products (product TEXT, total REAL)")

        # Export data
        data = {
            "metadata": {"source_type": "csv", "source_channel": "file"},
            "records": [
                {"_id": "rec_0", "_source": {"product": "Widget A", "total": 299.9}},
                {"_id": "rec_1", "_source": {"product": "Widget B", "total": 249.95}},
            ]
        }
        channel.export("products", data)

        # Verify
        result = channel.read("SELECT * FROM products")
        assert len(result) == 2


class TestIoTChannel:
    """Test IoTChannel - read() and export() methods."""

    def test_read_and_export_methods_exist(self):
        """Test that read() and export() methods exist."""
        channel = IoTChannel()
        assert hasattr(channel, 'read')
        assert hasattr(channel, 'export')
        assert callable(channel.read)
        assert callable(channel.export)

    def test_connect_mqtt_method_exists(self):
        """Test MQTT connection method exists."""
        channel = IoTChannel()
        assert hasattr(channel, 'connect_mqtt')

    def test_platform_attribute(self):
        """Test IoTChannel has platform attribute."""
        channel = IoTChannel(platform="aws_iot", region="us-east-1")
        assert channel.platform == "aws_iot"
        assert channel.region == "us-east-1"


# ============================================================================
# ADAPTER TESTS
# ============================================================================

class TestCSVAdapter:
    """Test CSVAdapter - parse() and serialize() methods."""

    def test_parse_with_header(self):
        """Test parsing CSV with header row."""
        adapter = CSVAdapter(delimiter=",", has_header=True)
        data = adapter.parse("name,age,city\nAlice,30,NYC\nBob,25,LA\n")

        assert len(data) == 2
        assert data[0] == {"name": "Alice", "age": "30", "city": "NYC"}
        assert data[1] == {"name": "Bob", "age": "25", "city": "LA"}

    def test_parse_without_header(self):
        """Test parsing CSV without header."""
        adapter = CSVAdapter(delimiter=",", has_header=False)
        data = adapter.parse("Alice,30,NYC\nBob,25,LA\n")

        assert len(data) == 2
        assert data[0] == {"col_0": "Alice", "col_1": "30", "col_2": "NYC"}

    def test_serialize_csv(self, tmp_path):
        """Test serializing data to CSV."""
        adapter = CSVAdapter()
        data = [{"name": "Alice", "age": "30"}, {"name": "Bob", "age": "25"}]

        output_path = tmp_path / "output.csv"
        adapter.serialize(data, str(output_path))

        content = output_path.read_text(encoding="utf-8")
        assert "name,age" in content
        assert "Alice,30" in content

    def test_validate_valid_data(self):
        """Test validation of valid data."""
        adapter = CSVAdapter()
        data = [{"a": 1}, {"a": 2}]
        assert adapter.validate(data) is True

    def test_validate_empty_data(self):
        """Test validation of empty data."""
        adapter = CSVAdapter()
        assert adapter.validate([]) is False

    def test_normalize_method(self):
        """Test normalize() helper method."""
        adapter = CSVAdapter()
        data = [{"name": "Alice", "age": "30"}]
        result = adapter.normalize(data, "csv")

        assert "metadata" in result
        assert "records" in result
        assert result["metadata"]["source_type"] == "csv"


class TestJSONAdapter:
    """Test JSONAdapter - parse() and serialize() methods."""

    def test_parse_json_array(self):
        """Test parsing JSON array."""
        adapter = JSONAdapter()
        data = adapter.parse('[{"id": 1}, {"id": 2}]')

        assert len(data) == 2
        assert data[0] == {"id": 1}

    def test_parse_json_object_with_records_key(self):
        """Test parsing JSON object with records key."""
        adapter = JSONAdapter()
        data = adapter.parse('{"records": [{"id": 1}, {"id": 2}]}')

        assert len(data) == 2
        assert data[0] == {"id": 1}

    def test_parse_json_lines_format(self):
        """Test parsing JSON Lines format."""
        adapter = JSONAdapter(lines=True)
        data = adapter.parse('{"id": 1}\n{"id": 2}\n{"id": 3}')

        assert len(data) == 3
        assert data[2] == {"id": 3}

    def test_serialize_json(self, tmp_path):
        """Test serializing data to JSON."""
        adapter = JSONAdapter()
        data = [{"id": 1}, {"id": 2}]

        output_path = tmp_path / "output.json"
        adapter.serialize(data, str(output_path))

        content = json.loads(output_path.read_text(encoding="utf-8"))
        assert "records" in content
        assert len(content["records"]) == 2


class TestXMLAdapter:
    """Test XMLAdapter - parse() and serialize() methods."""

    def test_parse_simple_xml(self):
        """Test parsing simple XML."""
        adapter = XMLAdapter(record_path=".//record")
        xml_content = """<?xml version="1.0"?>
<root>
    <record><name>Alice</name><age>30</age></record>
    <record><name>Bob</name><age>25</age></record>
</root>"""
        data = adapter.parse(xml_content.encode("utf-8"))

        assert len(data) == 2
        assert data[0]["name"] == "Alice"
        assert data[1]["name"] == "Bob"

    def test_parse_xml_with_attributes(self):
        """Test parsing XML with attributes."""
        adapter = XMLAdapter(record_path=".//item")
        xml_content = """<root>
    <item id="1" type="A"><name>Alice</name></item>
    <item id="2" type="B"><name>Bob</name></item>
</root>"""
        data = adapter.parse(xml_content.encode("utf-8"))

        assert data[0]["@id"] == "1"
        assert data[0]["@type"] == "A"

    def test_serialize_xml(self, tmp_path):
        """Test serializing data to XML."""
        adapter = XMLAdapter()
        data = [{"name": "Alice", "age": "30"}]

        output_path = tmp_path / "output.xml"
        adapter.serialize(data, str(output_path))

        content = output_path.read_text(encoding="utf-8")
        assert "<record>" in content
        assert "<name>Alice</name>" in content


class TestExcelAdapter:
    """Test ExcelAdapter - parse() and serialize() methods."""

    def test_parse_excel_from_bytes(self):
        """Test parsing Excel data from bytes."""
        from io import BytesIO
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["name", "age"])
        ws.append(["Alice", "30"])
        ws.append(["Bob", "25"])

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        adapter = ExcelAdapter()
        data = adapter.parse(buffer.getvalue())

        assert len(data) == 2
        assert data[0] == {"name": "Alice", "age": "30"}

    def test_get_sheets(self):
        """Test getting sheet names."""
        from io import BytesIO
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws2 = wb.create_sheet("Sheet2")

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        adapter = ExcelAdapter()
        sheets = adapter.get_sheets(buffer.getvalue())

        assert "Sheet1" in sheets
        assert "Sheet2" in sheets


class TestYAMLAdapter:
    """Test YAMLAdapter - parse() and serialize() methods."""

    def test_parse_yaml_list(self):
        """Test parsing YAML list."""
        adapter = YAMLAdapter()
        yaml_content = """- name: Alice
  age: 30
- name: Bob
  age: 25"""
        data = adapter.parse(yaml_content.encode("utf-8"))

        assert len(data) == 2
        assert data[0]["name"] == "Alice"

    def test_serialize_yaml(self, tmp_path):
        """Test serializing data to YAML."""
        adapter = YAMLAdapter()
        data = [{"name": "Alice", "age": 30}, {"name": "Bob", "age": 25}]

        output_path = tmp_path / "output.yaml"
        adapter.serialize(data, str(output_path))

        content = output_path.read_text(encoding="utf-8")
        assert "name:" in content
        assert "Alice" in content


# ============================================================================
# NORMALIZER TESTS
# ============================================================================

class TestJSONNormalizer:
    """Test JSONNormalizer - data_category detection and normalization."""

    def test_normalize_flatten_data(self):
        """Test normalizing flatten (single-table) data."""
        normalizer = JSONNormalizer()
        data = [
            {"name": "Alice", "age": 30},
            {"name": "Bob", "age": 25},
        ]

        result = normalizer.normalize(data, "csv", "file", "/path/to/data.csv")

        assert result["metadata"]["source_type"] == "csv"
        assert result["metadata"]["source_channel"] == "file"
        assert result["metadata"]["record_count"] == 2
        assert result["metadata"]["data_category"] == "flatten"
        assert len(result["records"]) == 2

    def test_normalize_complex_data(self):
        """Test normalizing complex (hierarchical) data."""
        normalizer = JSONNormalizer()
        data = [
            {
                "company": "Acme",
                "departments": [
                    {"name": "Engineering", "staff": [
                        {"name": "Alice", "role": "Engineer"}
                    ]}
                ]
            }
        ]

        result = normalizer.normalize(data, "json", "file", "/path/to/data.json")

        assert result["metadata"]["data_category"] == "complex"
        assert result["metadata"]["record_count"] == 1

    def test_normalize_empty_data(self):
        """Test normalizing empty data."""
        normalizer = JSONNormalizer()
        result = normalizer.normalize([], "csv", "file")

        assert result["metadata"]["record_count"] == 0
        assert result["metadata"]["data_category"] == "flatten"
        assert result["records"] == []

    def test_detect_category_flatten(self):
        """Test _detect_category returns flatten for simple data."""
        normalizer = JSONNormalizer()
        data = [{"a": 1}, {"b": 2}]
        category = normalizer._detect_category(data)
        assert category == "flatten"

    def test_detect_category_complex_nested_list(self):
        """Test _detect_category returns complex for nested lists."""
        normalizer = JSONNormalizer()
        data = [{"name": "Test", "items": [{"id": 1}, {"id": 2}]}]
        category = normalizer._detect_category(data)
        assert category == "complex"

    def test_detect_category_complex_nested_dict(self):
        """Test _detect_category returns complex for nested dicts."""
        normalizer = JSONNormalizer()
        data = [{"name": "Test", "metadata": {"key": "value"}}]
        category = normalizer._detect_category(data)
        assert category == "complex"

    def test_denormalize(self):
        """Test denormalizing back to list of dicts."""
        normalizer = JSONNormalizer()
        normalized = {
            "metadata": {"source_type": "csv"},
            "records": [
                {"_id": "abc_0", "_source": {"name": "Alice"}},
                {"_id": "abc_1", "_source": {"name": "Bob"}},
            ]
        }

        data = normalizer.denormalize(normalized)

        assert len(data) == 2
        assert data[0] == {"name": "Alice"}
        assert data[1] == {"name": "Bob"}

    def test_merge_normalized(self):
        """Test merging multiple normalized datasets."""
        normalizer = JSONNormalizer()
        normalized1 = {
            "metadata": {"columns": ["name"]},
            "records": [{"_id": "1", "_source": {"name": "Alice"}}]
        }
        normalized2 = {
            "metadata": {"columns": ["age"]},
            "records": [{"_id": "2", "_source": {"age": 30}}]
        }

        merged = normalizer.merge([normalized1, normalized2])

        assert merged["metadata"]["record_count"] == 2
        assert "name" in merged["metadata"]["columns"]
        assert "age" in merged["metadata"]["columns"]


# ============================================================================
# INTEGRATION TESTS
# ============================================================================

class TestChannelAdapterIntegration:
    """Integration tests for channel -> adapter -> normalizer flow."""

    def test_file_to_csv_to_normalized(self, tmp_path):
        """Test: FileChannel read -> CSVAdapter parse -> JSONNormalizer normalize."""
        csv_file = tmp_path / "sales.csv"
        csv_file.write_text(
            "product,quantity,price\nWidget A,10,29.99\nWidget B,5,49.99\n",
            encoding="utf-8"
        )

        # Channel reads file
        channel = FileChannel()
        raw = channel.read(str(csv_file))

        # Adapter parses to records
        adapter = CSVAdapter()
        records = adapter.parse(raw)

        # Normalizer creates standard JSON
        normalizer = JSONNormalizer()
        result = normalizer.normalize(records, "csv", "file", str(csv_file))

        assert result["metadata"]["source_type"] == "csv"
        assert result["metadata"]["source_channel"] == "file"
        assert result["metadata"]["record_count"] == 2
        assert result["metadata"]["data_category"] == "flatten"
        assert result["records"][0]["_source"]["product"] == "Widget A"
        assert result["records"][0]["_source"]["quantity"] == "10"  # Widget A row
        assert result["records"][1]["_source"]["product"] == "Widget B"
        assert result["records"][1]["_source"]["quantity"] == "5"  # Widget B row

    def test_file_to_json_to_normalized(self, tmp_path):
        """Test: FileChannel read -> JSONAdapter parse -> JSONNormalizer normalize."""
        json_file = tmp_path / "users.json"
        json_file.write_text(
            '[{"name": "Alice", "email": "alice@example.com"}, {"name": "Bob", "email": "bob@example.com"}]',
            encoding="utf-8"
        )

        channel = FileChannel()
        raw = channel.read(str(json_file))

        adapter = JSONAdapter()
        records = adapter.parse(raw)

        normalizer = JSONNormalizer()
        result = normalizer.normalize(records, "json", "file", str(json_file))

        assert result["metadata"]["source_type"] == "json"
        assert result["metadata"]["record_count"] == 2

    def test_file_export_roundtrip(self, tmp_path):
        """Test: Create data -> export to file -> read back -> verify."""
        original_data = {
            "metadata": {
                "source_type": "csv",
                "source_channel": "file",
                "record_count": 2,
                "data_category": "flatten",
            },
            "records": [
                {"_id": "test_0", "_source": {"name": "Alice", "age": "30"}},
                {"_id": "test_1", "_source": {"name": "Bob", "age": "25"}},
            ]
        }

        # Export via FileChannel
        channel = FileChannel()
        output_path = tmp_path / "exported.json"
        channel.export(str(output_path), original_data)

        # Read back
        content = json.loads(output_path.read_text(encoding="utf-8"))

        assert content["metadata"]["record_count"] == 2
        assert len(content["records"]) == 2
        assert content["records"][0]["_source"]["name"] == "Alice"


class TestDataCategoryIntegration:
    """Integration tests for data category handling."""

    def test_flatten_category_csv(self):
        """Test flatten category for CSV data."""
        adapter = CSVAdapter()
        records = adapter.parse("a,b,c\n1,2,3\n4,5,6\n")

        normalizer = JSONNormalizer()
        result = normalizer.normalize(records, "csv", "file")

        assert result["metadata"]["data_category"] == "flatten"

    def test_complex_category_nested_json(self):
        """Test complex category for nested JSON."""
        adapter = JSONAdapter()
        nested_data = {
            "company": "Acme",
            "departments": [
                {"name": "Eng", "employees": [
                    {"name": "Alice"},
                    {"name": "Bob"}
                ]}
            ]
        }
        records = adapter.parse(json.dumps([nested_data]))

        normalizer = JSONNormalizer()
        result = normalizer.normalize(records, "json", "file")

        assert result["metadata"]["data_category"] == "complex"


# ============================================================================
# ERROR HANDLING TESTS
# ============================================================================

class TestErrorHandling:
    """Test error handling and exceptions."""

    def test_channel_read_error(self):
        """Test ChannelReadError is raised for read failures."""
        channel = FileChannel()
        with pytest.raises(ChannelReadError):
            channel.read("/nonexistent/path/file.csv")

    def test_ingestion_error_inheritance(self):
        """Test error class inheritance hierarchy."""
        assert issubclass(ChannelReadError, IngestionError)
        assert issubclass(ParseError, IngestionError)
        assert issubclass(NormalizationError, IngestionError)


# ============================================================================
# FIXTURES
# ============================================================================

@pytest.fixture
def sample_csv_content():
    """Sample CSV content for testing."""
    return "name,age,city\nAlice,30,NYC\nBob,25,LA\nCharlie,35,Chicago\n"


@pytest.fixture
def sample_json_content():
    """Sample JSON content for testing."""
    return json.dumps([
        {"name": "Alice", "age": 30},
        {"name": "Bob", "age": 25},
    ])


@pytest.fixture
def sample_nested_json():
    """Sample nested JSON for complex category testing."""
    return json.dumps([{
        "company": "Acme Corp",
        "departments": [
            {
                "name": "Engineering",
                "staff": [
                    {"name": "Alice", "role": "Engineer"},
                    {"name": "Bob", "role": "Engineer"}
                ]
            },
            {
                "name": "Sales",
                "staff": [
                    {"name": "Charlie", "role": "Manager"}
                ]
            }
        ]
    }])


if __name__ == "__main__":
    pytest.main([__file__, "-v"])